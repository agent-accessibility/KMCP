"""Process uploaded documentation files into markdown-like text."""

from __future__ import annotations

import io
import logging
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from pptx import Presentation

logger = logging.getLogger(__name__)

_OCR_MIN_CHARS = 20


def _read_pdf(file_bytes: bytes) -> str:
    pages: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text = page.get_text("text").strip()
            if len(text) < _OCR_MIN_CHARS:
                try:
                    tp = page.get_textpage_ocr(full=True)
                    text = page.get_text("text", textpage=tp).strip()
                    if text:
                        logger.info("OCR fallback extracted %d chars on page %d", len(text), page.number)
                except RuntimeError:
                    logger.warning("OCR unavailable (Tesseract not installed?), skipping page %d", page.number)
            if text:
                pages.append(text)
    return "\n\n".join(pages)


def _read_docx(file_bytes: bytes) -> str:
    document = DocxDocument(io.BytesIO(file_bytes))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def _read_pptx(file_bytes: bytes) -> str:
    presentation = Presentation(io.BytesIO(file_bytes))
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
    return "\n".join(lines)


def process_file(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Convert uploaded files into text suitable for chunking."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        content = _read_pdf(file_bytes)
    elif suffix == ".docx":
        content = _read_docx(file_bytes)
    elif suffix == ".pptx":
        content = _read_pptx(file_bytes)
    else:
        content = file_bytes.decode("utf-8", errors="ignore")

    return content.strip(), {"title": Path(filename).stem, "filename": filename}